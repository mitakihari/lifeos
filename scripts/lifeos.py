#!/usr/bin/env python3
"""lifeos: command-line helper for a LifeOS Obsidian vault.

One file, standard library only for the core. Finance commands additionally
need `beancount` (books) and `openpyxl` (spreadsheet reports); `fava` is the
optional web dashboard.

Run `python3 lifeos.py --help` for the command list. Inside a vault the copy at
`_system/bin/lifeos.py` is used; inside the toolkit repo it is
`scripts/lifeos.py`. Both are the same file.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.error
import urllib.request
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:  # Python 3.11+
    import tomllib
except ModuleNotFoundError:  # pragma: no cover
    tomllib = None  # type: ignore

VERSION = "0.1.0"
MARKER = ".lifeos.json"

# Folders that hold "system" content rather than notes. Lint and other note
# commands skip them.
SYSTEM_DIRS = {".git", ".obsidian", ".claude", ".trash", "_system", "_templates", "_attachments", "finance", "node_modules"}
# Note folders whose notes are allowed to have no backlinks.
ORPHAN_OK_PREFIXES = ("50 Journal", "00 Inbox", "40 Archive", "_bases")

# Obsidian community plugins installed by `plugins install`: (plugin id, GitHub repo)
OBSIDIAN_PLUGINS = [
    ("obsidian-tasks-plugin", "obsidian-tasks-group/obsidian-tasks"),
    ("obsidian-git", "Vinzent03/obsidian-git"),
    ("calendar", "liamcain/obsidian-calendar-plugin"),
]

# Skills vendored from kepano/obsidian-skills that get synced into the vault.
VENDORED_SKILLS = ["obsidian-markdown", "obsidian-bases", "json-canvas"]


# --------------------------------------------------------------------------- #
# Paths and helpers
# --------------------------------------------------------------------------- #

def die(msg: str, code: int = 1) -> None:
    print(f"error: {msg}", file=sys.stderr)
    sys.exit(code)


def toolkit_root() -> Path | None:
    """The toolkit repo root, if this script is running from one."""
    here = Path(__file__).resolve().parent
    root = here.parent
    if (root / "template").is_dir() and (root / "claude").is_dir():
        return root
    return None


def find_vault(explicit: str | None) -> Path:
    """Vault = --vault, $LIFEOS_VAULT, or the nearest ancestor with .lifeos.json."""
    if explicit:
        p = Path(explicit).expanduser().resolve()
        if not (p / MARKER).exists():
            die(f"{p} is not a LifeOS vault (no {MARKER}). Run `init` first.")
        return p
    env = os.environ.get("LIFEOS_VAULT")
    if env:
        return find_vault(env)
    for start in (Path.cwd(), Path(__file__).resolve().parent):
        for p in [start, *start.parents]:
            if (p / MARKER).exists():
                return p
    die("could not find a LifeOS vault. Pass --vault PATH or run from inside the vault.")
    raise SystemExit  # for type checkers


def today() -> dt.date:
    override = os.environ.get("LIFEOS_TODAY")  # used by tests
    return dt.date.fromisoformat(override) if override else dt.date.today()


def iter_notes(vault: Path):
    for path in sorted(vault.rglob("*.md")):
        rel = path.relative_to(vault)
        if rel.parts and rel.parts[0] in SYSTEM_DIRS:
            continue
        if any(part.startswith(".") for part in rel.parts):
            continue
        yield path


FRONTMATTER_RE = re.compile(r"\A\ufeff?---\r?\n(.*?)\r?\n---\r?\n?", re.S)


def parse_frontmatter(text: str) -> dict[str, str]:
    """Tiny flat-YAML reader: `key: value` lines only. Good enough for our notes."""
    m = FRONTMATTER_RE.match(text)
    if not m:
        return {}
    out: dict[str, str] = {}
    for line in m.group(1).splitlines():
        if not line or line.startswith((" ", "\t", "-", "#")) or ":" not in line:
            continue
        line = line.rstrip("\r")
        key, _, value = line.partition(":")
        out[key.strip()] = value.strip().strip('"').strip("'")
    return out


# --------------------------------------------------------------------------- #
# init / sync / doctor
# --------------------------------------------------------------------------- #

def copy_tree(src: Path, dst: Path, overwrite: bool) -> list[Path]:
    written = []
    for item in sorted(src.rglob("*")):
        if item.is_dir() or item.name == ".DS_Store":
            continue
        target = dst / item.relative_to(src)
        if target.exists() and not overwrite:
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(item, target)
        written.append(target)
    return written


def cmd_init(args) -> None:
    root = toolkit_root()
    if not root:
        die("`init` must be run from the toolkit repo (scripts/lifeos.py).")
    vault = Path(args.path).expanduser().resolve()
    if vault.exists() and any(vault.iterdir()) and not args.adopt:
        die(f"{vault} is not empty. Use --adopt to add LifeOS to an existing vault "
            "(it only adds missing files and never overwrites).")
    vault.mkdir(parents=True, exist_ok=True)
    written = copy_tree(root / "template", vault, overwrite=False)
    # personalise
    name = args.name or "Me"
    for path in written:
        if path.suffix in {".md", ".beancount", ".toml"}:
            text = path.read_text(encoding="utf-8")
            new = text.replace("{{owner}}", name).replace("{{init_date}}", today().isoformat())
            if new != text:
                path.write_text(new, encoding="utf-8")
    print(f"Created {len(written)} files in {vault}")
    sync(vault, root)
    py = "python" if os.name == "nt" else "python3"
    print("\nNext: open this folder as a vault in Obsidian, then run\n"
          f"  {py} {vault / '_system/bin/lifeos.py'} plugins install\n"
          f"  {py} {vault / '_system/bin/lifeos.py'} doctor")


def sync(vault: Path, root: Path) -> None:
    """Copy managed system files (Claude commands/skills, CLI, hooks) into the vault."""
    managed: list[Path] = []
    claude_dir = vault / ".claude"
    managed += copy_tree(root / "claude" / "commands", claude_dir / "commands", overwrite=True)
    managed += copy_tree(root / "claude" / "skills", claude_dir / "skills", overwrite=True)
    vendor = root / "vendor" / "obsidian-skills"
    for skill in VENDORED_SKILLS:
        src = vendor / "skills" / skill
        if src.is_dir():
            managed += copy_tree(src, claude_dir / "skills" / skill, overwrite=True)
    if (vendor / "LICENSE").exists():
        lic = claude_dir / "skills" / "OBSIDIAN-SKILLS-LICENSE.txt"
        shutil.copy2(vendor / "LICENSE", lic)
        managed.append(lic)
    bin_dir = vault / "_system" / "bin"
    bin_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(Path(__file__).resolve(), bin_dir / "lifeos.py")
    managed.append(bin_dir / "lifeos.py")
    hooks_dir = vault / "_system" / "hooks"
    managed += copy_tree(root / "hooks", hooks_dir, overwrite=True)
    for hook in hooks_dir.iterdir():
        hook.chmod(0o755)
    marker = vault / MARKER
    data = json.loads(marker.read_text()) if marker.exists() else {"created": today().isoformat()}
    data.update({"toolkit_version": VERSION, "last_sync": today().isoformat()})
    marker.write_text(json.dumps(data, indent=2) + "\n")
    if (vault / ".git").exists():
        subprocess.run(["git", "-C", str(vault), "config", "core.hooksPath", "_system/hooks"], check=False)
    print(f"Synced {len(managed)} managed files into {vault} (toolkit v{VERSION})")


def cmd_sync(args) -> None:
    root = toolkit_root()
    if not root:
        die("`sync` must be run from the toolkit repo so it has something to copy from.")
    sync(find_vault(args.vault), root)


def cmd_doctor(args) -> None:
    vault = find_vault(args.vault)
    ok = True

    def check(label: str, passed: bool, fix: str = "") -> None:
        nonlocal ok
        ok &= passed
        print(f"  [{'ok' if passed else '!!'}] {label}" + ("" if passed or not fix else f"\n        fix: {fix}"))

    print(f"LifeOS vault: {vault}")
    check(f"Python {sys.version.split()[0]} (need 3.11+)", sys.version_info >= (3, 11), "install Python 3.11 or newer")
    for mod, why in [("beancount", "finance books"), ("openpyxl", "spreadsheet reports"), ("fava", "finance web dashboard")]:
        try:
            __import__(mod)
            check(f"{mod} installed ({why})", True)
        except ImportError:
            check(f"{mod} installed ({why})", False, f"pip install {mod}")
    check("git installed", shutil.which("git") is not None, "install Git (see docs/SETUP.md)")
    check("vault is a git repo", (vault / ".git").exists(), "see docs/SETUP.md, step 4")
    if (vault / ".git").exists():
        hp = subprocess.run(["git", "-C", str(vault), "config", "core.hooksPath"], capture_output=True, text=True).stdout.strip()
        check("secret-scan hook active", hp == "_system/hooks", "git config core.hooksPath _system/hooks")
    plugins_dir = vault / ".obsidian" / "plugins"
    for pid, _ in OBSIDIAN_PLUGINS:
        check(f"Obsidian plugin {pid}", (plugins_dir / pid / "main.js").exists(), "lifeos.py plugins install")
    check("CLAUDE.md present", (vault / "CLAUDE.md").exists())
    if (vault / "finance" / "main.beancount").exists():
        try:
            errors = load_ledger(vault)[1]
            check(f"ledger loads cleanly ({len(errors)} errors)", not errors, "lifeos.py finance check")
        except SystemExit:
            pass
    print("\nAll good." if ok else "\nSome checks need attention (see fixes above).")


# --------------------------------------------------------------------------- #
# new: create notes from templates
# --------------------------------------------------------------------------- #

NOTE_KINDS = {
    # kind: (template file, folder, filename pattern)
    "daily": ("Daily.md", "50 Journal/Daily", "{date}"),
    "weekly": ("Weekly Review.md", "50 Journal/Weekly", "{week}"),
    "monthly": ("Monthly Review.md", "50 Journal/Monthly", "{month}"),
    "project": ("Project.md", "10 Projects", "{title}"),
    "course": ("Course.md", "20 Areas/School/Courses", "{title}"),
    "application": ("Job Application.md", "20 Areas/Career/Applications", "{title}"),
    "listing": ("Apartment Listing.md", "20 Areas/Housing/Listings", "{title}"),
    "decision": ("Decision.md", "30 Resources/Decisions", "{date} {title}"),
    "resource": ("Resource.md", "30 Resources", "{title}"),
}

SAFE_NAME_RE = re.compile(r'[\\/:*?"<>|#^\[\]]')


def render(text: str, title: str, when: dt.date) -> str:
    iso = when.isocalendar()
    week_start = when - dt.timedelta(days=when.weekday())
    subs = {
        "{{title}}": title,
        "{{date}}": when.isoformat(),
        "{{time}}": dt.datetime.now().strftime("%H:%M"),
        "{{week}}": f"{iso[0]}-W{iso[1]:02d}",
        "{{week_start}}": week_start.isoformat(),
        "{{week_end}}": (week_start + dt.timedelta(days=6)).isoformat(),
        "{{month}}": when.strftime("%Y-%m"),
        "{{month_name}}": when.strftime("%B %Y"),
        "{{yesterday}}": (when - dt.timedelta(days=1)).isoformat(),
        "{{tomorrow}}": (when + dt.timedelta(days=1)).isoformat(),
    }
    for k, v in subs.items():
        text = text.replace(k, v)
    return text


def cmd_new(args) -> None:
    vault = find_vault(args.vault)
    kind = args.kind
    template, folder, pattern = NOTE_KINDS[kind]
    when = dt.date.fromisoformat(args.date) if args.date else today()
    title = " ".join(args.title).strip()
    if "{title}" in pattern and not title:
        die(f"`new {kind}` needs a title, e.g. lifeos.py new {kind} \"Something\"")
    title = SAFE_NAME_RE.sub("-", title)
    iso = when.isocalendar()
    name = pattern.format(date=when.isoformat(), week=f"{iso[0]}-W{iso[1]:02d}",
                          month=when.strftime("%Y-%m"), title=title).strip()
    target = vault / folder / f"{name}.md"
    if target.exists():
        print(target.relative_to(vault).as_posix())
        return
    src = vault / "_templates" / template
    if not src.exists():
        die(f"template missing: {src.relative_to(vault)}")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(render(src.read_text(encoding="utf-8"), title or name, when), encoding="utf-8")
    print(target.relative_to(vault).as_posix())


# --------------------------------------------------------------------------- #
# lint
# --------------------------------------------------------------------------- #

WIKILINK_RE = re.compile(r"(?<!`)!?\[\[([^\]|#^]+)(?:[#^][^\]|]*)?(?:\|[^\]]*)?\]\]")
TASK_RE = re.compile(r"^\s*[-*] \[( |x|X|/|-)\] (.*)$")
DUE_RE = re.compile(r"📅 ?(\d{4}-\d{2}-\d{2})")
CODE_FENCE_RE = re.compile(r"```.*?```", re.S)
INLINE_CODE_RE = re.compile(r"`[^`\n]*`")
JOURNAL_NAME_RE = re.compile(r"^\d{4}-(\d{2}-\d{2}|W\d{2}|\d{2})$")  # daily/weekly/monthly nav links
# Note types that are listed by Bases views or hubs, so they don't need backlinks.
INDEXED_TYPES = {"project", "course", "application", "listing", "daily", "weekly", "monthly", "area"}
DATE_FIELDS = ("due", "next_step_date", "available")
DONE_STATUSES = {"done", "dropped", "rejected", "withdrawn", "ghosted", "signed", "archived", "completed", "cancelled"}


def link_targets(vault: Path) -> set[str]:
    names = set()
    for p in vault.rglob("*"):
        rel = p.relative_to(vault)
        if p.is_file() and not any(part.startswith(".") for part in rel.parts):
            names.add(p.stem.lower())
            names.add(p.name.lower())
            names.add(str(rel.with_suffix("")).lower().replace(os.sep, "/"))
            names.add(str(rel).lower().replace(os.sep, "/"))
    return names


def collect_lint(vault: Path) -> dict:
    targets = link_targets(vault)
    notes = list(iter_notes(vault))
    backlinks: dict[str, int] = defaultdict(int)
    types: dict[str, str] = {}
    findings = {"dead_links": [], "orphans": [], "missing_type": [], "overdue": [], "due_soon": [], "inbox": []}
    now = today()
    soon = now + dt.timedelta(days=7)
    for path in notes:
        rel = str(path.relative_to(vault)).replace(os.sep, "/")
        text = path.read_text(encoding="utf-8", errors="replace")
        scan = INLINE_CODE_RE.sub("", CODE_FENCE_RE.sub("", text))
        for m in WIKILINK_RE.finditer(scan):
            target = m.group(1).strip()
            key = target.lower()
            backlinks[key.split("/")[-1]] += 1
            if key not in targets and key + ".md" not in targets and not JOURNAL_NAME_RE.match(target):
                findings["dead_links"].append({"note": rel, "link": target})
        fm = parse_frontmatter(text)
        types[rel] = fm.get("type", "")
        if rel.startswith("00 Inbox/"):
            findings["inbox"].append(rel)
        elif rel.count("/") >= 1 and not rel.startswith(("_bases",)) and "type" not in fm:
            findings["missing_type"].append(rel)
        status = fm.get("status", "").lower()
        for field in DATE_FIELDS:
            val = fm.get(field, "")
            try:
                d = dt.date.fromisoformat(val[:10])
            except ValueError:
                continue
            if status in DONE_STATUSES:
                continue
            item = {"note": rel, "what": field, "date": d.isoformat()}
            if d < now:
                findings["overdue"].append(item)
            elif d <= soon:
                findings["due_soon"].append(item)
        for line in scan.splitlines():
            tm = TASK_RE.match(line)
            if not tm or tm.group(1) in "xX-":
                continue
            dm = DUE_RE.search(tm.group(2))
            if not dm:
                continue
            d = dt.date.fromisoformat(dm.group(1))
            item = {"note": rel, "what": tm.group(2).strip()[:90], "date": d.isoformat()}
            if d < now:
                findings["overdue"].append(item)
            elif d <= soon:
                findings["due_soon"].append(item)
    for path in notes:
        rel = str(path.relative_to(vault)).replace(os.sep, "/")
        if rel.startswith(ORPHAN_OK_PREFIXES) or "/" not in rel:
            continue
        if types.get(rel) in INDEXED_TYPES:
            continue
        if backlinks.get(path.stem.lower(), 0) == 0:
            findings["orphans"].append(rel)
    for key in ("overdue", "due_soon"):
        findings[key].sort(key=lambda i: i["date"])
    return findings


def cmd_lint(args) -> None:
    vault = find_vault(args.vault)
    f = collect_lint(vault)
    if args.json:
        print(json.dumps(f, indent=2))
        return
    labels = [
        ("overdue", "Overdue"), ("due_soon", "Due in the next 7 days"), ("inbox", "Inbox items to file"),
        ("dead_links", "Dead links"), ("missing_type", "Notes missing a `type:` property"), ("orphans", "Orphan notes (nothing links to them)"),
    ]
    total = 0
    for key, label in labels:
        items = f[key]
        if not items:
            continue
        total += len(items)
        print(f"\n{label} ({len(items)})")
        for it in items:
            if isinstance(it, str):
                print(f"  - {it}")
            elif "link" in it:
                print(f"  - {it['note']} -> [[{it['link']}]]")
            else:
                print(f"  - {it['date']}  {it['what']}  ({it['note']})")
    print("\nVault is clean." if total == 0 else f"\n{total} findings.")
    if args.strict and (f["dead_links"] or f["missing_type"]):
        sys.exit(1)


# --------------------------------------------------------------------------- #
# secret scan (used by the pre-commit hook)
# --------------------------------------------------------------------------- #

CARD_RE = re.compile(r"(?<![\d.-])(?:\d{13,19}|\d{4}(?:[ -]\d{4}){2}[ -]\d{1,7})(?![\d.])")
SSN_RE = re.compile(r"(?<!\d)\d{3}-\d{2}-\d{4}(?!\d)")
ACCT_RE = re.compile(r"(?i)\b(?:account|acct|routing)\s*(?:number|no\.?|#)?\s*[:=]?\s*\d{6,}")
SECRET_KEY_RE = re.compile(r"(?i)\b(?:api[_-]?key|secret|password|passwd|token)\s*[:=]\s*\S{8,}")


def luhn(digits: str) -> bool:
    total, alt = 0, False
    for ch in reversed(digits):
        n = int(ch)
        if alt:
            n = n * 2 - 9 if n * 2 > 9 else n * 2
        total += n
        alt = not alt
    return total % 10 == 0


def scan_text(text: str) -> list[str]:
    hits = []
    for m in CARD_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group())
        if 13 <= len(digits) <= 19 and luhn(digits) and len(set(digits)) > 1:
            hits.append(f"possible card number ending {digits[-4:]}")
    if SSN_RE.search(text):
        hits.append("possible Social Security number")
    if ACCT_RE.search(text):
        hits.append("possible bank account/routing number")
    if SECRET_KEY_RE.search(text):
        hits.append("possible password/API key")
    return hits


def cmd_scan(args) -> None:
    problems: list[tuple[str, str]] = []
    if args.staged:
        files = subprocess.run(["git", "diff", "--cached", "--name-only", "--diff-filter=ACM"],
                               capture_output=True, text=True, check=True).stdout.splitlines()
        for f in files:
            if f.startswith("finance/imports/raw/") and not f.endswith(".gitkeep"):
                problems.append((f, "raw bank export (keep these out of git)"))
                continue
            blob = subprocess.run(["git", "show", f":{f}"], capture_output=True)
            if blob.returncode or b"\0" in blob.stdout[:4096]:
                continue
            for hit in scan_text(blob.stdout.decode("utf-8", errors="replace")):
                problems.append((f, hit))
    else:
        for p in args.paths:
            for hit in scan_text(Path(p).read_text(encoding="utf-8", errors="replace")):
                problems.append((p, hit))
    if problems:
        print("LifeOS secret scan blocked this commit:", file=sys.stderr)
        for f, hit in problems:
            print(f"  {f}: {hit}", file=sys.stderr)
        print("Remove the sensitive value (use the last 4 digits only), then commit again.\n"
              "If it's a false alarm: git commit --no-verify", file=sys.stderr)
        sys.exit(1)
    print("secret scan: clean")


# --------------------------------------------------------------------------- #
# finance: Beancount books, CSV import, spreadsheet + markdown reports
# --------------------------------------------------------------------------- #

def need(module: str):
    try:
        return __import__(module)
    except ImportError:
        die(f"this command needs `{module}`. Install it with: pip install {module}")


def load_ledger(vault: Path):
    need("beancount")
    from beancount import loader
    main = vault / "finance" / "main.beancount"
    if not main.exists():
        die(f"no ledger at {main}")
    return loader.load_file(str(main))


def cmd_finance_check(args) -> None:
    vault = find_vault(args.vault)
    entries, errors, _ = load_ledger(vault)
    from beancount.core import data
    txns = [e for e in entries if isinstance(e, data.Transaction)]
    for e in errors:
        src = getattr(e, "source", None) or {}
        loc = f"{Path(src.get('filename', '?')).name}:{src.get('lineno', '?')}"
        print(f"  {loc}  {e.message}")
    unc = sum(1 for t in txns for p in t.postings if p.account.endswith(":Uncategorized"))
    print(f"{len(txns)} transactions, {len(errors)} errors, {unc} uncategorized postings")
    if errors:
        sys.exit(1)


def load_rules(vault: Path) -> dict:
    if tomllib is None:
        die("Python 3.11+ is required for import rules (tomllib).")
    path = vault / "finance" / "import-rules.toml"
    if not path.exists():
        die(f"missing {path.relative_to(vault)}")
    with open(path, "rb") as fh:
        return tomllib.load(fh)


def parse_amount(raw: str) -> Decimal | None:
    s = (raw or "").strip().replace("$", "").replace(",", "")
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        val = Decimal(s)
    except InvalidOperation:
        return None
    return -val if neg else val


def bean_str(s: str) -> str:
    return s.replace("\\", "").replace('"', "'").strip()


def existing_import_ids(vault: Path) -> set[str]:
    ids = set()
    for p in (vault / "finance").rglob("*.beancount"):
        ids.update(re.findall(r'import-id:\s*"([0-9a-f]+)"', p.read_text(encoding="utf-8")))
    return ids


US_STATES = set("AL AK AZ AR CA CO CT DE FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY "
                 "NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY DC".split())


def clean_payee(desc: str) -> str:
    """'HANNAFORD #8240 AUBURN ME' -> 'Hannaford'. Best effort; rules can set an exact payee."""
    s = re.sub(r"(?i)\b(pos|debit|purchase|card|recurring|ach|dbt|pmt)\b", " ", desc)
    s = re.split(r"\s#\s?\d+|\s\d{3,}", s)[0]
    words = s.split()
    if len(words) > 2 and words[-1].upper() in US_STATES:
        words = words[:-2]  # drop "<CITY> <ST>"
    s = " ".join(words).strip(" -*")
    return s.title() if s else desc.title()


def categorize(desc: str, amount: Decimal, rules: list[dict]) -> tuple[str | None, str | None]:
    for r in rules:
        if r.get("direction") == "in" and amount < 0:
            continue
        if r.get("direction") == "out" and amount > 0:
            continue
        if re.search(r["match"], desc, re.I):
            return r["account"], r.get("payee")
    return None, None


def cmd_finance_import(args) -> None:
    vault = find_vault(args.vault)
    cfg = load_rules(vault)
    profiles = cfg.get("profiles", {})
    if args.profile not in profiles:
        die(f"unknown profile '{args.profile}'. Defined: {', '.join(profiles) or 'none'} (finance/import-rules.toml)")
    prof = profiles[args.profile]
    rules = cfg.get("rules", [])
    seen = existing_import_ids(vault)
    csv_path = Path(args.csv).expanduser()
    with open(csv_path, newline="", encoding=prof.get("encoding", "utf-8-sig")) as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        die("CSV has no rows")
    missing = [c for c in [prof["date"], prof["description"]] if c not in rows[0]]
    if missing:
        die(f"CSV is missing columns {missing}. Columns found: {list(rows[0])}")
    by_month: dict[str, list[str]] = defaultdict(list)
    counts = defaultdict(int)
    added = skipped = uncategorized = 0
    for row in rows:
        try:
            date = dt.datetime.strptime(row[prof["date"]].strip(), prof.get("date_format", "%m/%d/%Y")).date()
        except ValueError:
            continue  # summary/footer lines
        if "amount" in prof:
            amount = parse_amount(row.get(prof["amount"], ""))
        else:
            debit = parse_amount(row.get(prof.get("debit", ""), "")) or Decimal(0)
            credit = parse_amount(row.get(prof.get("credit", ""), "")) or Decimal(0)
            amount = credit - abs(debit)
        if amount is None or amount == 0:
            continue
        if prof.get("negate"):
            amount = -amount
        desc = " ".join(row[prof["description"]].split())
        key = (date.isoformat(), desc, str(amount))
        counts[key] += 1
        iid = hashlib.sha1("|".join([args.profile, *key, str(counts[key])]).encode()).hexdigest()[:16]
        if iid in seen:
            skipped += 1
            continue
        account, payee = categorize(desc, amount, rules)
        if not account:
            account = "Income:Uncategorized" if amount > 0 else "Expenses:Uncategorized"
            uncategorized += 1
        flag = "!" if account.endswith("Uncategorized") else "*"
        by_month[date.strftime("%Y-%m")].append(
            f'{date.isoformat()} {flag} "{bean_str(payee or clean_payee(desc))}" "{bean_str(desc)}"\n'
            f'  import-id: "{iid}"\n'
            f"  {prof['account']}  {amount:.2f} USD\n"
            f"  {account}\n"
        )
        added += 1
    out_dir = vault / "finance" / "transactions"
    out_dir.mkdir(parents=True, exist_ok=True)
    if not args.dry_run:
        for month, txns in sorted(by_month.items()):
            path = out_dir / f"{month}.beancount"
            header = "" if path.exists() else f"; Transactions for {month}. Imported and hand-entered.\n\n"
            with open(path, "a", encoding="utf-8") as fh:
                fh.write(header + "\n".join(sorted(txns)) + "\n")
    print(f"{'Would add' if args.dry_run else 'Added'} {added} transactions "
          f"({uncategorized} need a category, flagged with !), skipped {skipped} already imported.")
    if uncategorized:
        print("Fix them in finance/transactions/*.beancount, or add a rule to finance/import-rules.toml and re-run.")


def money(d: Decimal) -> float:
    return float(round(d, 2))


def ledger_facts(vault: Path, year: int) -> dict:
    """Pull everything the reports need out of the ledger in one pass."""
    entries, errors, _ = load_ledger(vault)
    from beancount.core import data
    rows = []          # income/expense postings for the year
    balances: dict[str, Decimal] = defaultdict(Decimal)
    opened = []
    for e in entries:
        if isinstance(e, data.Open):
            opened.append(e.account)
        if not isinstance(e, data.Transaction):
            continue
        for p in e.postings:
            if p.units is None or p.units.currency != "USD":
                continue
            amt = p.units.number
            if p.account.startswith(("Assets", "Liabilities")) and e.date <= today():
                balances[p.account] += amt
            if e.date.year != year or not p.account.startswith(("Income", "Expenses")):
                continue
            kind = "Income" if p.account.startswith("Income") else "Expense"
            rows.append({
                "date": e.date, "month": e.date.strftime("%Y-%m"), "payee": e.payee or "",
                "narration": e.narration or "", "account": p.account,
                "category": p.account.split(":", 1)[1], "type": kind,
                "amount": -amt if kind == "Income" else amt,
            })
    rows.sort(key=lambda r: (r["date"], r["account"]))
    return {"rows": rows, "balances": balances, "opened": opened, "errors": errors}


def load_budget(vault: Path) -> dict[str, Decimal]:
    path = vault / "finance" / "budget.toml"
    if not path.exists() or tomllib is None:
        return {}
    with open(path, "rb") as fh:
        raw = tomllib.load(fh).get("monthly", {})
    return {k.split(":", 1)[1] if k.startswith(("Expenses:", "Income:")) else k: Decimal(str(v)) for k, v in raw.items()}


def build_workbook(facts: dict, budget: dict[str, Decimal], year: int, out: Path) -> None:
    need("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    title_f = Font(name=F, size=14, bold=True)
    head_f = Font(name=F, size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="305496")
    input_f = Font(name=F, size=10, color="0000FF")
    input_fill = PatternFill("solid", start_color="FFFF00")
    note_f = Font(name=F, size=9, italic=True, color="595959")
    total_border = Border(top=Side(style="thin"))
    USD = '$#,##0.00;($#,##0.00);"-"'
    PCT = '0.0%;(0.0%);"-"'

    rows = facts["rows"]
    months = [f"{year}-{m:02d}" for m in range(1, 13)]
    exp_cats = sorted({r["category"] for r in rows if r["type"] == "Expense"} |
                      set(budget) |
                      {a.split(":", 1)[1] for a in facts["opened"] if a.startswith("Expenses:")})
    inc_cats = sorted({r["category"] for r in rows if r["type"] == "Income"} |
                      {a.split(":", 1)[1] for a in facts["opened"] if a.startswith("Income:")})

    wb = Workbook()

    def header(ws, row, labels, widths=None):
        for i, label in enumerate(labels, 1):
            c = ws.cell(row=row, column=i, value=label)
            c.font, c.fill = head_f, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w

    # ---- Read Me -----------------------------------------------------------
    ws = wb.active
    ws.title = "Read Me"
    ws["A1"], ws["A1"].font = f"Finance Report {year}", title_f
    lines = [
        f"Generated {today().isoformat()} from finance/main.beancount by `lifeos.py finance report`.",
        "",
        "Tabs:",
        "  Summary: this month and year-to-date at a glance, plus what you own and owe.",
        "  Monthly: spending and income by category for every month, against your budget.",
        "  Budget: your monthly targets. The yellow cells with blue text are the only cells meant to be edited.",
        "  Transactions: every income and expense line from the ledger for the year.",
        "",
        "How to change things:",
        "  Budget numbers: edit finance/budget.toml (or the Budget tab, but the next report rebuilds it from budget.toml).",
        "  Transactions: edit the ledger in finance/transactions/, never this file. This file is rebuilt every run.",
        "",
        "Color key: blue text on yellow = input you can change; black = formula; green = pulled from another tab.",
    ]
    for i, line in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=line).font = base
    ws.column_dimensions["A"].width = 110

    # ---- Transactions ------------------------------------------------------
    tx = wb.create_sheet("Transactions")
    header(tx, 1, ["Date", "Month", "Payee", "Description", "Account", "Category", "Type", "Amount ($)"],
           [12, 10, 28, 40, 34, 28, 10, 14])
    for i, r in enumerate(rows, 2):
        vals = [r["date"], r["month"], r["payee"], r["narration"], r["account"], r["category"], r["type"], money(r["amount"])]
        for j, v in enumerate(vals, 1):
            c = tx.cell(row=i, column=j, value=v)
            c.font = base
        tx.cell(row=i, column=1).number_format = "yyyy-mm-dd"
        tx.cell(row=i, column=8).number_format = USD
    last_tx = max(len(rows) + 1, 2)
    tx.freeze_panes = "A2"
    if rows:
        tx.auto_filter.ref = f"A1:H{last_tx}"
    note_row = last_tx + 2
    tx.cell(row=note_row, column=1, value="Source: finance/main.beancount (ledger). Income shown as positive, expenses as positive spending.").font = note_f
    rng = lambda col: f"Transactions!${col}$2:${col}${last_tx}"  # noqa: E731

    # ---- Budget ------------------------------------------------------------
    bd = wb.create_sheet("Budget")
    header(bd, 1, ["Category", "Monthly budget ($)"], [34, 18])
    budget_row: dict[str, int] = {}
    for i, cat in enumerate(exp_cats, 2):
        bd.cell(row=i, column=1, value=cat).font = base
        c = bd.cell(row=i, column=2, value=money(budget.get(cat, Decimal(0))))
        c.font, c.fill, c.number_format = input_f, input_fill, USD
        budget_row[cat] = i
    btot = len(exp_cats) + 2
    bd.cell(row=btot, column=1, value="Total monthly budget").font = bold
    c = bd.cell(row=btot, column=2, value=f"=SUM(B2:B{btot - 1})")
    c.font, c.number_format, c.border = bold, USD, total_border
    bd.cell(row=btot + 2, column=1, value="Source: finance/budget.toml (amounts you set). 0 means no budget set yet.").font = note_f

    # ---- Monthly -----------------------------------------------------------
    mo = wb.create_sheet("Monthly")
    mlabels = [dt.date(year, m, 1).strftime("%b") for m in range(1, 13)]
    header(mo, 1, ["Category", "Budget / mo", *mlabels, "YTD total", "Avg / mo"], [30, 12, *[11] * 12, 13, 12])
    months_row = 2
    mo.cell(row=months_row, column=1, value="Month key").font = note_f
    for j, mk in enumerate(months, 3):
        c = mo.cell(row=months_row, column=j, value=mk)
        c.font, c.alignment = note_f, Alignment(horizontal="center")
    cur_month_idx = today().month if today().year == year else 12
    elapsed = max(1, cur_month_idx)

    def section(start: int, title: str, cats: list[str], kind: str, with_budget: bool) -> tuple[int, int]:
        mo.cell(row=start, column=1, value=title).font = bold
        first = start + 1
        for i, cat in enumerate(cats, first):
            mo.cell(row=i, column=1, value=cat).font = base
            if with_budget and cat in budget_row:
                c = mo.cell(row=i, column=2, value=f"=Budget!B{budget_row[cat]}")
                c.font = Font(name=F, size=10, color="008000")
                c.number_format = USD
            for j in range(3, 15):
                col = get_column_letter(j)
                c = mo.cell(row=i, column=j,
                            value=f'=SUMIFS({rng("H")},{rng("F")},$A{i},{rng("B")},{col}${months_row},{rng("G")},"{kind}")')
                c.font, c.number_format = base, USD
            c = mo.cell(row=i, column=15, value=f"=SUM(C{i}:N{i})")
            c.font, c.number_format = bold, USD
            c = mo.cell(row=i, column=16, value=f"=O{i}/{elapsed}")
            c.font, c.number_format = base, USD
        last = first + len(cats) - 1 if cats else first
        tot = last + 1
        mo.cell(row=tot, column=1, value=f"Total {title.lower()}").font = bold
        for j in range(2, 17):
            col = get_column_letter(j)
            c = mo.cell(row=tot, column=j, value=f"=SUM({col}{first}:{col}{last})" if cats else 0)
            c.font, c.number_format, c.border = bold, USD, total_border
        return first, tot

    _, inc_tot = section(4, "Income", inc_cats, "Income", with_budget=False)
    _, exp_tot = section(inc_tot + 2, "Expenses", exp_cats, "Expense", with_budget=True)
    net = exp_tot + 2
    mo.cell(row=net, column=1, value="Net saved (income - expenses)").font = bold
    mo.cell(row=net + 1, column=1, value="Savings rate").font = bold
    for j in range(3, 17):
        col = get_column_letter(j)
        c = mo.cell(row=net, column=j, value=f"={col}{inc_tot}-{col}{exp_tot}")
        c.font, c.number_format = bold, USD
        c = mo.cell(row=net + 1, column=j, value=f"=IF({col}{inc_tot}=0,0,{col}{net}/{col}{inc_tot})")
        c.font, c.number_format = base, PCT
    mo.cell(row=net + 3, column=1,
            value=f"Avg / mo divides YTD by {elapsed} month(s) elapsed in {year}. Budget pulled from the Budget tab.").font = note_f
    mo.freeze_panes = "C3"

    # ---- Summary -----------------------------------------------------------
    sm = wb.create_sheet("Summary", 1)
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 16
    sm.column_dimensions["C"].width = 16
    sm["A1"], sm["A1"].font = f"Money Summary {year}", title_f
    cur_col = get_column_letter(2 + cur_month_idx)
    header(sm, 3, ["", dt.date(year, cur_month_idx, 1).strftime("%B"), "Year to date"])
    items = [
        ("Income", inc_tot), ("Spending", exp_tot), ("Net saved", net),
    ]
    for i, (label, r) in enumerate(items, 4):
        sm.cell(row=i, column=1, value=label).font = bold
        for col, src in ((2, cur_col), (3, "O")):
            c = sm.cell(row=i, column=col, value=f"=Monthly!{src}{r}")
            c.font, c.number_format = Font(name=F, size=10, color="008000"), USD
    sm.cell(row=7, column=1, value="Savings rate").font = bold
    for col in ("B", "C"):
        c = sm[f"{col}7"]
        c.value = f"=IF({col}4=0,0,{col}6/{col}4)"
        c.font, c.number_format = base, PCT
    sm.cell(row=8, column=1, value="Budget this month").font = bold
    c = sm.cell(row=8, column=2, value=f"=Budget!B{btot}")
    c.font, c.number_format = Font(name=F, size=10, color="008000"), USD
    sm.cell(row=9, column=1, value="Left in budget this month").font = bold
    c = sm.cell(row=9, column=2, value="=B8-B5")
    c.font, c.number_format = bold, USD

    header(sm, 11, ["Account", f"Balance as of {today().isoformat()}", ""])
    r = 12
    accts = sorted(a for a in facts["opened"] if a.startswith(("Assets", "Liabilities")))
    first_bal = r
    for a in accts:
        sm.cell(row=r, column=1, value=a).font = base
        c = sm.cell(row=r, column=2, value=money(facts["balances"].get(a, Decimal(0))))
        c.font, c.number_format = base, USD
        r += 1
    sm.cell(row=r, column=1, value="Net worth (own - owe)").font = bold
    c = sm.cell(row=r, column=2, value=f"=SUM(B{first_bal}:B{max(first_bal, r - 1)})")
    c.font, c.number_format, c.border = bold, USD, total_border
    sm.cell(row=r + 2, column=1,
            value="Balances are ledger values (debts show negative). They are only right if opening balances and all transactions are entered.").font = note_f

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = sheet.title in {"Transactions", "Monthly", "Budget"}
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def markdown_summary(facts: dict, budget: dict[str, Decimal], year: int) -> str:
    rows = facts["rows"]
    month = today().strftime("%Y-%m") if today().year == year else f"{year}-12"
    def total(kind, m=None):
        return sum((r["amount"] for r in rows if r["type"] == kind and (m is None or r["month"] == m)), Decimal(0))
    inc_m, exp_m = total("Income", month), total("Expense", month)
    inc_y, exp_y = total("Income"), total("Expense")
    by_cat: dict[str, Decimal] = defaultdict(Decimal)
    for r in rows:
        if r["type"] == "Expense" and r["month"] == month:
            by_cat[r["category"]] += r["amount"]
    unc = sum(1 for r in rows if r["category"].startswith("Uncategorized"))
    lines = [
        f"_Updated {today().isoformat()} by `lifeos.py finance report`. Numbers come from the ledger; don't edit this block._",
        "",
        f"| | {dt.date.fromisoformat(month + '-01').strftime('%B')} | Year to date |",
        "|---|---:|---:|",
        f"| Income | ${inc_m:,.2f} | ${inc_y:,.2f} |",
        f"| Spending | ${exp_m:,.2f} | ${exp_y:,.2f} |",
        f"| Net saved | ${inc_m - exp_m:,.2f} | ${inc_y - exp_y:,.2f} |",
        "",
        "**This month by category vs budget**",
        "",
        "| Category | Spent | Budget | Left |",
        "|---|---:|---:|---:|",
    ]
    cats = sorted(set(by_cat) | {c for c, v in budget.items() if v})
    for c in cats:
        b = budget.get(c, Decimal(0))
        left = f"${b - by_cat[c]:,.2f}" if b else "-"
        lines.append(f"| {c} | ${by_cat[c]:,.2f} | {'$%s' % format(b, ',.2f') if b else '-'} | {left} |")
    if not cats:
        lines.append("| (no spending yet) | | | |")
    lines += ["", "**Balances**", "", "| Account | Balance |", "|---|---:|"]
    nw = Decimal(0)
    for a in sorted(x for x in facts["opened"] if x.startswith(("Assets", "Liabilities"))):
        bal = facts["balances"].get(a, Decimal(0))
        nw += bal
        lines.append(f"| {a} | ${bal:,.2f} |")
    lines.append(f"| **Net worth** | **${nw:,.2f}** |")
    if unc:
        lines += ["", f"> [!warning] {unc} transaction line(s) this year still need a category."]
    return "\n".join(lines)


def replace_block(text: str, name: str, body: str) -> str:
    start, end = f"<!-- lifeos:{name}:start -->", f"<!-- lifeos:{name}:end -->"
    block = f"{start}\n{body}\n{end}"
    if start in text and end in text:
        return re.sub(re.escape(start) + r".*?" + re.escape(end), lambda _: block, text, flags=re.S)
    return text.rstrip() + "\n\n" + block + "\n"


def cmd_finance_report(args) -> None:
    vault = find_vault(args.vault)
    year = args.year or today().year
    facts = ledger_facts(vault, year)
    if facts["errors"]:
        print(f"warning: ledger has {len(facts['errors'])} error(s); run `finance check`.", file=sys.stderr)
    budget = load_budget(vault)
    out = vault / "finance" / "reports" / f"Finance {year}.xlsx"
    build_workbook(facts, budget, year, out)
    note = vault / "20 Areas" / "Money" / "Money.md"
    if note.exists():
        note.write_text(replace_block(note.read_text(encoding="utf-8"), "finance", markdown_summary(facts, budget, year)),
                        encoding="utf-8")
    print(f"Wrote {out.relative_to(vault)}" + (f" and updated {note.relative_to(vault)}" if note.exists() else ""))


def cmd_finance_fava(args) -> None:
    vault = find_vault(args.vault)
    need("fava")
    print("Fava running at http://localhost:5000 (Ctrl+C to stop)")
    # Run through this Python rather than a `fava` executable, which pip --user
    # often puts in a folder that isn't on PATH (especially on Windows).
    subprocess.run([sys.executable, "-c", "from fava.cli import main; main()",
                    str(vault / "finance" / "main.beancount")])


# --------------------------------------------------------------------------- #
# Obsidian plugins
# --------------------------------------------------------------------------- #

def fetch(url: str) -> bytes | None:
    req = urllib.request.Request(url, headers={"User-Agent": "lifeos"})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.read()
    except urllib.error.HTTPError as e:
        if e.code == 404:
            return None
        raise


def cmd_plugins_install(args) -> None:
    vault = find_vault(args.vault)
    obs = vault / ".obsidian"
    plugins_dir = obs / "plugins"
    installed = []
    for pid, repo in OBSIDIAN_PLUGINS:
        base = f"https://github.com/{repo}/releases/latest/download/"
        try:
            manifest = fetch(base + "manifest.json")
            main = fetch(base + "main.js")
        except (urllib.error.URLError, TimeoutError) as e:
            print(f"  !! {pid}: download failed ({e}). Install it from Obsidian > Settings > Community plugins.")
            continue
        if not manifest or not main:
            print(f"  !! {pid}: release files not found on {repo}")
            continue
        real_id = json.loads(manifest).get("id", pid)
        dest = plugins_dir / real_id
        dest.mkdir(parents=True, exist_ok=True)
        (dest / "manifest.json").write_bytes(manifest)
        (dest / "main.js").write_bytes(main)
        styles = fetch(base + "styles.css")
        if styles:
            (dest / "styles.css").write_bytes(styles)
        installed.append(real_id)
        print(f"  ok {real_id} {json.loads(manifest).get('version', '')}")
    cp = obs / "community-plugins.json"
    enabled = json.loads(cp.read_text()) if cp.exists() else []
    for pid in installed:
        if pid not in enabled:
            enabled.append(pid)
    obs.mkdir(exist_ok=True)
    cp.write_text(json.dumps(enabled, indent=2) + "\n")
    print("Restart Obsidian (or reload the vault). If it asks, turn off Restricted mode to enable community plugins.")


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def main(argv: list[str] | None = None) -> None:
    for stream in (sys.stdout, sys.stderr):  # Windows consoles default to cp1252, which can't print emoji
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    p = argparse.ArgumentParser(prog="lifeos", description="LifeOS vault helper")
    p.add_argument("--version", action="version", version=f"lifeos {VERSION}")
    p.add_argument("--vault", help="path to the vault (default: auto-detect)")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("init", help="create a new vault from the template (toolkit only)")
    s.add_argument("path")
    s.add_argument("--name", help="your first name, used in notes")
    s.add_argument("--adopt", action="store_true", help="add LifeOS to an existing, non-empty folder")
    s.set_defaults(func=cmd_init)

    sub.add_parser("sync", help="update the vault's Claude commands, skills, CLI and hooks (toolkit only)").set_defaults(func=cmd_sync)
    sub.add_parser("doctor", help="check that everything is installed and wired up").set_defaults(func=cmd_doctor)

    s = sub.add_parser("new", help="create a note from a template")
    s.add_argument("kind", choices=sorted(NOTE_KINDS))
    s.add_argument("title", nargs="*")
    s.add_argument("--date", help="YYYY-MM-DD (default today)")
    s.set_defaults(func=cmd_new)

    s = sub.add_parser("lint", help="find overdue items, dead links, orphans, unfiled inbox notes")
    s.add_argument("--json", action="store_true")
    s.add_argument("--strict", action="store_true", help="exit 1 on dead links or missing types")
    s.set_defaults(func=cmd_lint)

    s = sub.add_parser("scan", help="check files for card/account numbers and secrets")
    s.add_argument("--staged", action="store_true", help="scan files staged for git commit")
    s.add_argument("paths", nargs="*")
    s.set_defaults(func=cmd_scan)

    fin = sub.add_parser("finance", help="books, imports and reports").add_subparsers(dest="fcmd", required=True)
    fin.add_parser("check", help="validate the ledger").set_defaults(func=cmd_finance_check)
    s = fin.add_parser("import", help="import a bank/card CSV export")
    s.add_argument("csv")
    s.add_argument("--profile", required=True, help="profile name from finance/import-rules.toml")
    s.add_argument("--dry-run", action="store_true")
    s.set_defaults(func=cmd_finance_import)
    s = fin.add_parser("report", help="build the yearly spreadsheet and refresh Money.md")
    s.add_argument("--year", type=int)
    s.set_defaults(func=cmd_finance_report)
    fin.add_parser("fava", help="open the Fava web dashboard").set_defaults(func=cmd_finance_fava)

    pl = sub.add_parser("plugins", help="Obsidian community plugins").add_subparsers(dest="pcmd", required=True)
    pl.add_parser("install", help="download and enable the LifeOS plugin set").set_defaults(func=cmd_plugins_install)

    # Accept --vault anywhere on the line, not only before the command.
    argv = list(sys.argv[1:] if argv is None else argv)
    vault = None
    for i, a in enumerate(argv):
        if a == "--vault" and i + 1 < len(argv):
            vault = argv[i + 1]
            del argv[i:i + 2]
            break
        if a.startswith("--vault="):
            vault = a.split("=", 1)[1]
            del argv[i]
            break
    args = p.parse_args(argv)
    if vault:
        args.vault = vault
    args.func(args)


if __name__ == "__main__":
    main()
